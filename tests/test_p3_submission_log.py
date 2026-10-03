"""P3-16: submission log schema, migration, private-limit and XLSX tests."""
import pytest

from src.p3.submission_log import (
    LOG_COLUMNS,
    PRIVATE_LIMIT,
    append_submission,
    export_xlsx,
    read_log,
    summarize_log,
    write_log,
)


def record(run_id="run1", split="public", **overrides):
    row = {"run_id": run_id, "split": split, "archive": f"submission_{run_id}.zip",
           "git_commit": "a" * 40, "config": "configs/p3/official_submission.yaml",
           "model": "qwen3-reranker", "local_score": 0.42, "public_score": 0.4,
           "private_score": "", "notes": ""}
    row.update(overrides)
    return row


def test_read_log_missing_file_returns_empty(tmp_path):
    assert read_log(tmp_path / "missing.csv") == []


def test_append_creates_p3_header_and_row(tmp_path):
    path = tmp_path / "submission_log.csv"
    append_submission(path, record())
    rows = read_log(path)
    assert len(rows) == 1 and rows[0]["run_id"] == "run1"
    with path.open(encoding="utf-8") as handle:
        assert handle.readline().strip() == ",".join(LOG_COLUMNS)


def test_legacy_score_column_is_migrated(tmp_path):
    path = tmp_path / "submission_log.csv"
    path.write_text("run_id,split,archive,score,notes\nlegacy1,public,sub.zip,0.31,old row\n",
                    encoding="utf-8")
    rows = read_log(path)
    assert rows[0]["local_score"] == "0.31" and rows[0]["notes"] == "old row"
    write_log(path, rows)
    with path.open(encoding="utf-8") as handle:
        assert handle.readline().strip() == ",".join(LOG_COLUMNS)
    assert read_log(path)[0]["local_score"] == "0.31"


def test_private_limit_blocks_sixth_private_run(tmp_path):
    path = tmp_path / "submission_log.csv"
    for index in range(PRIVATE_LIMIT):
        append_submission(path, record(run_id=f"priv{index}", split="private"))
    with pytest.raises(ValueError):
        append_submission(path, record(run_id="priv6", split="private"))
    append_submission(path, record(run_id="priv6", split="private"), enforce_private_limit=False)
    append_submission(path, record(run_id="pub1", split="public"))
    summary = summarize_log(read_log(path))
    assert summary["private_count"] == PRIVATE_LIMIT + 1
    assert summary["public_count"] == 1 and summary["private_remaining"] == 0


def test_duplicate_run_and_invalid_split_rejected(tmp_path):
    path = tmp_path / "submission_log.csv"
    append_submission(path, record())
    with pytest.raises(ValueError):
        append_submission(path, record())
    with pytest.raises(ValueError):
        append_submission(path, {"run_id": "x", "split": "bogus"})


def test_summarize_reports_best_local_score(tmp_path):
    path = tmp_path / "submission_log.csv"
    append_submission(path, record(run_id="r1", local_score=0.2))
    append_submission(path, record(run_id="r2", local_score=0.8, split="private"))
    summary = summarize_log(read_log(path))
    assert summary["best_local_score"] == pytest.approx(0.8)
    assert summary["total"] == 2


def test_export_xlsx_is_exclusive(tmp_path):
    pytest.importorskip("openpyxl")
    path = tmp_path / "submission_log.xlsx"
    rows = [record()]
    assert export_xlsx(path, rows).endswith("submission_log.xlsx")
    from openpyxl import load_workbook
    sheet = load_workbook(path).active
    assert [cell.value for cell in sheet[1]] == list(LOG_COLUMNS)
    assert sheet.max_row == 2
    with pytest.raises(FileExistsError):
        export_xlsx(path, rows)
